"""生成 12 页答辩 PPT（python-pptx，可编辑文本/表格/图片）。

用法（使用 DSH 自带 Python，已含 python-pptx / Pillow）：
    python scripts/make_ppt.py
输出：
    reports/airline_satisfaction_ppt.pptx

设计口径与 reports/ppt_p2.md 一致；数字来自 reports/exp_log.csv 与 reports/eval/。
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "reports" / "figures"
OUT = ROOT / "reports" / "airline_satisfaction_ppt.pptx"

# ---------------------------------------------------------------- 主题
SW, SH = 13.333, 7.5
M = 0.62
CW = SW - 2 * M
NAVY = RGBColor(0x0F, 0x27, 0x40)
NAVY_LIGHT = RGBColor(0x16, 0x3A, 0x5A)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
TEAL_LIGHT = RGBColor(0xE6, 0xF4, 0xF5)
AMBER = RGBColor(0xD9, 0x77, 0x06)
AMBER_LIGHT = RGBColor(0xFF, 0xF7, 0xED)
INK = RGBColor(0x1F, 0x29, 0x37)
SLATE = RGBColor(0x47, 0x55, 0x69)
MUTED = RGBColor(0x64, 0x74, 0x8B)
BORDER = RGBColor(0xD9, 0xE2, 0xEC)
LIGHT = RGBColor(0xF4, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xB9, 0x1C, 0x1C)
GREEN = RGBColor(0x04, 0x78, 0x57)
FONT = "Microsoft YaHei"


def _font(run, size: float, bold: bool = False, color: RGBColor = INK,
          name: str = FONT, italic: bool = False) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = rpr.makeelement(qn("a:ea"), {})
        rpr.append(ea)
    ea.set("typeface", name)
    cs = rpr.find(qn("a:cs"))
    if cs is None:
        cs = rpr.makeelement(qn("a:cs"), {})
        rpr.append(cs)
    cs.set("typeface", name)


def text(slide, x, y, w, h, content, size=14, bold=False, color=INK,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line=1.15,
         space_after=0, margins=(0.0, 0.0, 0.0, 0.0)):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left, tf.margin_right, tf.margin_top, tf.margin_bottom = [Inches(v) for v in margins]
    tf.vertical_anchor = anchor
    for i, line_text in enumerate(str(content).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = line_text
        _font(r, size, bold, color)
    return box


def rich(slide, x, y, w, h, paragraphs, anchor=MSO_ANCHOR.TOP, margins=(0, 0, 0, 0)):
    """paragraphs: list of dict(text|runs, size, bold, color, align, space_after, line)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left, tf.margin_right, tf.margin_top, tf.margin_bottom = [Inches(v) for v in margins]
    tf.vertical_anchor = anchor
    for i, spec in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = spec.get("align", PP_ALIGN.LEFT)
        p.line_spacing = spec.get("line", 1.12)
        p.space_after = Pt(spec.get("space_after", 0))
        p.space_before = Pt(spec.get("space_before", 0))
        runs = spec.get("runs") or [(spec.get("text", ""), {})]
        for txt, style in runs:
            r = p.add_run()
            r.text = str(txt)
            _font(
                r,
                style.get("size", spec.get("size", 14)),
                style.get("bold", spec.get("bold", False)),
                style.get("color", spec.get("color", INK)),
                style.get("name", FONT),
                style.get("italic", False),
            )
    return box


def rect(slide, x, y, w, h, fill, line=None, rounded=False, radius=0.06, line_w=0.75):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    if rounded:
        shape.adjustments[0] = radius
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_w)
    shape.shadow.inherit = False
    return shape


def accent(slide, x, y, w, h, color=TEAL):
    rect(slide, x, y, w, h, color)


def card(slide, x, y, w, h, fill=LIGHT, line=BORDER, bar=None):
    rect(slide, x, y, w, h, fill, line, rounded=True, radius=0.05)
    if bar:
        accent(slide, x, y + 0.12, 0.055, h - 0.24, bar)


def pic(slide, path: Path, x, y, w, h, border=False):
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    left, top = x + (w - pw) / 2, y + (h - ph) / 2
    if border:
        rect(slide, left - 0.02, top - 0.02, pw + 0.04, ph + 0.04, WHITE, BORDER, line_w=0.5)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(pw), Inches(ph))


def header(slide, kicker, title, idx):
    rect(slide, 0, 0, 0.14, SH, NAVY)
    rect(slide, 0, 0, 0.14, 0.72, TEAL)
    text(slide, M, 0.25, 9.8, 0.28, kicker, 10.5, True, TEAL)
    text(slide, M, 0.52, CW, 0.68, title, 26, True, NAVY)
    rect(slide, M, 1.32, CW, 0.012, BORDER)
    # footer
    rect(slide, M, 7.08, CW, 0.008, BORDER)
    text(slide, M, 7.14, 8.4, 0.24, "Kaggle S6E10 · 航空乘客满意度预测", 9, False, MUTED)
    text(slide, SW - M - 0.8, 7.14, 0.8, 0.24, f"{idx:02d}", 9, True, MUTED, PP_ALIGN.RIGHT)


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body


def add_cover(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY
    # decorative circles
    c1 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.4), Inches(-1.6), Inches(4.6), Inches(4.6))
    c1.fill.solid(); c1.fill.fore_color.rgb = NAVY_LIGHT; c1.line.fill.background(); c1.shadow.inherit = False
    c2 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(11.8), Inches(4.6), Inches(3.4), Inches(3.4))
    c2.fill.solid(); c2.fill.fore_color.rgb = RGBColor(0x13, 0x32, 0x4F); c2.line.fill.background(); c2.shadow.inherit = False
    rect(s, 0, 0, 0.18, SH, TEAL)
    text(s, 0.9, 0.95, 7.6, 0.3, "KAGGLE PLAYGROUND SERIES · S6E10", 12, True, TEAL_LIGHT)
    text(s, 0.9, 1.45, 8.2, 1.0, "航空乘客满意度预测", 44, True, WHITE)
    text(s, 0.9, 2.55, 7.8, 0.55, "从数据、特征到模型与评估的全流程协作", 21, False, RGBColor(0xCB, 0xD5, 0xE1))
    rect(s, 0.92, 3.45, 1.1, 0.05, AMBER)
    text(s, 0.9, 3.75, 7.5, 0.9,
         "二分类 · ROC AUC · 提交概率\nLightGBM + base+C+T · 5 折 OOF 0.9602",
         15, False, RGBColor(0xC7, 0xD2, 0xDF), line=1.45)
    text(s, 0.9, 6.35, 7.8, 0.35, "P1 数据与特征  ·  P2 模型与训练  ·  P3 评估与交付", 11.5, True, RGBColor(0x94, 0xA3, 0xB8))
    text(s, 0.9, 6.78, 6.5, 0.3, "2026.10 · 课程项目答辩", 10.5, False, RGBColor(0x94, 0xA3, 0xB8))

    # stats panel
    rect(s, 9.05, 1.35, 3.35, 4.85, RGBColor(0x14, 0x33, 0x50), RGBColor(0x2A, 0x5A, 0x7A), rounded=True, radius=0.04)
    stats = [
        ("0.9602", "OOF ROC AUC", TEAL_LIGHT),
        ("699,635", "训练样本 × 23 列", WHITE),
        ("13", "项 0–5 服务评分", WHITE),
    ]
    yy = 1.72
    for i, (num, label, color) in enumerate(stats):
        text(s, 9.35, yy, 2.8, 0.62, num, 27, True, color)
        text(s, 9.35, yy + 0.62, 2.8, 0.3, label, 10, False, RGBColor(0xB6, 0xC4, 0xD4))
        yy += 1.42
        if i < 2:
            rect(s, 9.35, yy - 0.22, 2.65, 0.008, RGBColor(0x2A, 0x5A, 0x7A))
    text(s, 9.35, 5.75, 2.75, 0.3, "提交概率 · 不阈值化", 9.5, True, AMBER)
    notes(s, "开场：一句话说明题目和交付——用乘客画像、行程与 13 项服务评分预测满意度概率，"
             "主指标 ROC AUC；最终 5 折 OOF 0.9602。强调这是三人协作、全流程可复现的课程项目。")
    return s


def add_slide2(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "01 / 背景与问题", "为什么做航空满意度预测", 2)
    items = [
        ("01", "业务价值", "满意度影响复购、口碑与客户终身价值；若能在飞行前识别高风险乘客，可在值机、登机、客舱服务环节提前干预。", TEAL),
        ("02", "任务定义", "二分类：输入乘客画像、行程信息与 13 项服务评分，输出 satisfaction 的概率；主指标 ROC AUC。", AMBER),
        ("03", "评测约束", "提交概率、禁止阈值化；固定 StratifiedKFold 5 折；所有实验必须可复现、可追溯。", NAVY),
    ]
    y = 1.62
    for num, title, body, color in items:
        card(s, M, y, 7.15, 1.42, WHITE, BORDER, color)
        text(s, M + 0.22, y + 0.2, 0.5, 0.4, num, 12, True, color)
        text(s, M + 0.82, y + 0.17, 6.0, 0.32, title, 15, True, NAVY)
        text(s, M + 0.82, y + 0.55, 6.05, 0.75, body, 11.8, False, SLATE, line=1.25)
        y += 1.55

    # right panel
    card(s, 8.05, 1.62, 4.66, 4.82, LIGHT, BORDER, TEAL)
    text(s, 8.32, 1.82, 4.1, 0.3, "核心问题", 11, True, TEAL)
    rich(s, 8.32, 2.18, 4.1, 1.25, [
        {"text": "哪些服务环节真正驱动满意度？", "size": 15, "bold": True, "color": NAVY, "space_after": 8, "line": 1.2},
        {"text": "只用飞行前可获得的信息，模型能做到多准？", "size": 15, "bold": True, "color": NAVY, "line": 1.2},
    ])
    rect(s, 8.32, 3.58, 4.05, 0.01, BORDER)
    text(s, 8.32, 3.75, 4.1, 0.3, "评价标准", 11, True, TEAL)
    text(s, 8.32, 4.08, 4.1, 1.25,
         "不是刷到最高分，而是：\n可复现 · 可解释 · 可部署",
         14, True, NAVY, line=1.5)
    text(s, 8.32, 5.55, 4.1, 0.7,
         "最终分数 0.9602 已接近公开分析天花板，\n项目重心转向评估与可解释性。",
         10.5, False, SLATE, line=1.3)

    rect(s, M, 6.55, CW, 0.38, NAVY)
    text(s, M + 0.2, 6.61, CW - 0.4, 0.26,
         "一句话：用可复现的实验回答“哪些服务信号真正重要”，而不是只报一个分数。",
         11.5, True, WHITE)
    notes(s, "背景讲业务价值：满意度与复购/口碑相关；飞行前识别高风险乘客可提前干预。"
             "任务定义为二分类、AUC、提交概率；评测约束是固定折、可追溯。右侧点出两个核心问题，"
             "底部强调项目评价标准不是单纯刷分。")
    return s


def add_slide3(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "02 / 数据与 EDA", "数据质量与一个反直觉事实：0 分不是“最差”", 3)

    # left quality
    text(s, M, 1.55, 5.4, 0.3, "数据体检", 11, True, TEAL)
    q = [
        ("缺失", "仅 Arrival Delay：train 292 / test 130（0.04%）"),
        ("重复", "按除 id 外全列判重：0 行"),
        ("结构", "无时间列、无实体列；train/test 无漂移"),
    ]
    yy = 1.9
    for key, body in q:
        text(s, M, yy, 0.62, 0.3, key, 10.5, True, MUTED)
        text(s, M + 0.7, yy, 4.65, 0.3, body, 11, False, SLATE)
        yy += 0.42

    # highlight card
    card(s, M, 3.35, 5.55, 2.55, AMBER_LIGHT, RGBColor(0xF5, 0xC9, 0x8A), AMBER)
    text(s, M + 0.22, 3.52, 5.1, 0.32, "0 分 = “不适用”，不是“最差”", 12.5, True, AMBER)
    rows = [
        ("Inflight wifi service", "89%", "39%"),
        ("Online boarding", "62%", "10.5%"),
    ]
    text(s, M + 0.22, 3.98, 2.5, 0.25, "评分列", 9.5, True, MUTED)
    text(s, M + 3.05, 3.98, 0.9, 0.25, "0 分", 9.5, True, MUTED, PP_ALIGN.CENTER)
    text(s, M + 4.15, 3.98, 1.0, 0.25, "1 分", 9.5, True, MUTED, PP_ALIGN.CENTER)
    yy = 4.25
    for name, v0, v1 in rows:
        text(s, M + 0.22, yy, 2.75, 0.3, name, 10.5, False, INK)
        text(s, M + 2.85, yy, 1.2, 0.32, v0, 15, True, GREEN, PP_ALIGN.CENTER)
        text(s, M + 3.95, yy, 1.0, 0.32, v1, 15, True, RED, PP_ALIGN.CENTER)
        yy += 0.42
    text(s, M + 0.22, 5.15, 5.05, 0.55,
         "0 分的满意率反而更高 → 直接决定特征组 A：\n0→NaN + is_zero 标记，不能把 0 当序数最低分。",
         10.5, False, SLATE, line=1.25)

    # right figure
    pic(s, FIG / "fig4_rating_vs_target.png", 6.4, 1.55, 6.3, 4.25, border=True)
    text(s, 6.4, 5.86, 6.3, 0.3, "评分–满意率：0 分处出现明显的非单调跳变（P1 EDA）", 9.5, False, MUTED, PP_ALIGN.CENTER)

    # bottom split
    rect(s, M, 6.22, CW, 0.55, LIGHT, BORDER, rounded=True, radius=0.15)
    chips = [("KFold", "0.95877"), ("Stratified（冻结）", "0.958829"), ("GroupKFold", "0.95778")]
    xx = M + 0.18
    for name, auc in chips:
        text(s, xx, 6.32, 2.0, 0.25, name, 9.5, True, MUTED)
        text(s, xx, 6.5, 0.9, 0.22, auc, 11, True, TEAL)
        xx += 2.05
    text(s, 7.1, 6.33, 5.45, 0.32,
         "按距离段分组只掉 0.001 → 模型学到的是可迁移服务信号，不是记住航线。",
         10.2, False, SLATE)
    notes(s, "数据体检三件事：只有到达延误少量缺失、无重复、无时间/实体列。"
             "重点讲 0 分语义：WiFi 0 分满意率 89% 而 1 分只有 39%，Online boarding 0 分 62% 对 1 分 10.5%，"
             "所以特征是 0→NaN + is_zero。右图是直接证据。底部说明三种划分，GroupKFold 只掉 0.001。")
    return s


def add_slide4(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "03 / 特征工程", "五组可开关特征：价值取决于模型能否表达", 4)

    text(s, M, 1.55, 4.7, 0.3, "特征组（全开 49 列）", 11, True, TEAL)
    data = [
        ["组", "内容", "列数", "用标签"],
        ["base", "原始 21 特征 + arr_delay_missing", "22", "否"],
        ["A", "0→NaN + is_zero", "13", "否"],
        ["B", "聚合 / 交互（隐含 A）", "11", "否"],
        ["C", "Flight Distance / Age 频次", "2", "否"],
        ["T", "Flight Distance 目标编码", "1", "是"],
        ["合计", "全开", "49", "—"],
    ]
    add_table(s, M, 1.9, 4.825, 3.55, data, [0.72, 2.55, 0.62, 0.75], 10.5,
              align=[PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.CENTER])

    text(s, 5.62, 1.55, 7.0, 0.3, "消融矩阵（5 折 OOF AUC，fold_v1，seed 42）", 11, True, TEAL)
    abl = [
        ["特征组", "维数", "LightGBM", "Δ", "逻辑回归", "Δ"],
        ["base", "22", "0.958829", "—", "0.921550", "—"],
        ["+A", "35", "0.958800", "−0.000029", "0.940807", "+0.019257"],
        ["+B", "46", "0.958650", "−0.000179", "0.943866", "+0.022316"],
        ["+C", "24", "0.959413", "+0.000584", "0.921643", "+0.000093"],
        ["+C+T", "25", "0.960191", "+0.001362", "0.923300", "+0.001750"],
        ["A+B+C+T", "49", "0.960048", "+0.001219", "0.945162", "+0.023612"],
    ]
    add_table(s, 5.62, 1.9, 7.05, 3.55, abl, [1.15, 0.55, 1.05, 1.0, 1.0, 1.0], 10.2,
              align=[PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.CENTER, PP_ALIGN.CENTER,
                     PP_ALIGN.CENTER, PP_ALIGN.CENTER],
              highlight_rows={5}, highlight_fill=TEAL_LIGHT)

    # insight cards
    yy, hh = 5.72, 1.0
    cards = [
        ("树模型", "A/B 零增益或负增益", "聚合量是原始列的确定性函数，树自己能学；多列反而稀释采样。", TEAL),
        ("线性模型", "A/B +0.019 ~ +0.022", "0 的非单调语义线性模型表达不了；is_zero 标记把它救回来。", AMBER),
        ("精选", "C+T 0.960191 > 全加 0.960048", "树模型唯一稳定增量是频次 + 目标编码；全加不如精选。", GREEN),
    ]
    x = M
    for tag, big, small, color in cards:
        card(s, x, yy, 3.92, hh, WHITE, BORDER, color)
        text(s, x + 0.18, yy + 0.1, 1.6, 0.25, tag, 9.5, True, color)
        text(s, x + 0.18, yy + 0.33, 3.55, 0.28, big, 12.5, True, NAVY)
        text(s, x + 0.18, yy + 0.62, 3.55, 0.34, small, 9.0, False, SLATE, line=1.15)
        x += 4.08

    text(s, M, 6.82, CW, 0.22,
         "注：B 组 dist_bin 已改为全量固定边界后重跑（run_id 以 _abl2 结尾）；主模型 base+C+T 不含 B，主分数不变。",
         9.0, False, MUTED)
    notes(s, "先讲五组特征和列数账本：base 22、A 13、B 11、C 2、T 1，全开 49。"
             "消融表是核心：同一组 A/B 对 LightGBM 是 0 或负，对 LR 是 +0.019~+0.022；"
             "树模型唯一稳定增量是 C+T +0.0014，而且全加 49 列 0.960048 反而不如精选 25 列 0.960191。"
             "结论一句话：特征的价值取决于模型能不能表达。")
    return s


def add_slide5(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "04 / 模型", "主力模型：LightGBM，线性模型作为对照", 5)

    # score comparison
    card(s, M, 1.58, 5.65, 1.95, WHITE, BORDER, TEAL)
    text(s, M + 0.25, 1.72, 2.0, 0.28, "LightGBM", 11, True, TEAL)
    text(s, M + 0.25, 2.0, 2.4, 0.7, "0.9602", 34, True, NAVY)
    text(s, M + 0.25, 2.72, 2.5, 0.28, "OOF AUC · base+C+T", 9.5, False, MUTED)
    rect(s, M + 2.92, 1.75, 0.008, 1.55, BORDER)
    text(s, M + 3.18, 1.72, 2.2, 0.28, "LogisticRegression", 11, True, AMBER)
    text(s, M + 3.18, 2.0, 2.3, 0.7, "0.9233", 34, True, NAVY)
    text(s, M + 3.18, 2.72, 2.4, 0.28, "OOF AUC · 同折同特征", 9.5, False, MUTED)
    text(s, M + 0.25, 3.08, 5.1, 0.3, "Δ = +0.0369，线性模型表达不了高阶交互", 10.5, True, RED)

    # reasons
    bullets = [
        ("非线性与交互", "舱位 × 旅程类型 × 延误的耦合，GBDT 天然拟合。"),
        ("类别特征友好", "直接处理 category，不用大规模 One-Hot。"),
        ("概率校准好", "Logloss 目标输出良好校准的概率，满足“提交概率”。"),
    ]
    yy = 3.62
    for i, (h, b) in enumerate(bullets, 1):
        text(s, M, yy, 0.3, 0.28, f"{i}", 10, True, TEAL)
        text(s, M + 0.34, yy - 0.02, 4.9, 0.28, h, 12.5, True, NAVY)
        text(s, M + 0.34, yy + 0.28, 4.95, 0.3, b, 10.2, False, SLATE)
        yy += 0.78

    # structure image
    pic(s, FIG / "p2_structure.png", 6.45, 1.6, 6.25, 2.9, border=True)
    text(s, 6.45, 4.55, 6.25, 0.3, "LightGBM：K 棵树 → 叶值加权和 → Sigmoid → 概率", 9.5, False, MUTED, PP_ALIGN.CENTER)

    # bottom band
    rect(s, M, 6.05, CW, 0.55, NAVY)
    text(s, M + 0.22, 6.13, CW - 0.44, 0.35,
         "选型不是拍脑袋：同折、同特征、只换模型，LightGBM 领先 0.0369。",
         12, True, WHITE)
    notes(s, "模型选型：LightGBM 0.9602 对 LR 0.9233，同折同特征只换模型，差 0.0369。"
             "三个理由：非线性交互、类别特征友好、Logloss 概率校准。右图是机制示意：K 棵树、叶值加权和、Sigmoid。"
             "底部强调这是一个受控对比，不是拍脑袋。")
    return s


def add_slide6(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "05 / 训练与调参", "训练协议与过拟合控制", 6)

    pic(s, FIG / "p2_learning_curve.png", M, 1.55, 6.45, 3.75, border=True)
    text(s, M, 5.38, 6.45, 0.3, "fold 0 学习曲线：train/valid AUC vs boosting round", 9.5, False, MUTED, PP_ALIGN.CENTER)

    stats = [
        ("5 折冻结", "StratifiedKFold · seed 42", TEAL),
        ("早停轮次", "best_iter ≈ 489", AMBER),
        ("train–valid 差距", "末轮约 0.0095，全程未发散", GREEN),
    ]
    yy = 1.55
    for big, small, color in stats:
        card(s, 7.3, yy, 5.35, 0.86, WHITE, BORDER, color)
        text(s, 7.48, yy + 0.14, 1.9, 0.3, big, 11, True, color)
        text(s, 7.48, yy + 0.47, 4.75, 0.26, small, 10.5, False, SLATE)
        yy += 0.98

    # sensitivity table
    text(s, 7.3, 4.62, 5.2, 0.3, "num_leaves 敏感性（其余默认）", 11, True, TEAL)
    data = [
        ["num_leaves", "OOF AUC", "相对 63"],
        ["31", "0.959895", "−0.000296"],
        ["63（默认）", "0.960191", "基准"],
        ["127", "0.960254", "+0.000063"],
    ]
    add_table(s, 7.3, 4.95, 5.3, 1.25, data, [1.45, 1.75, 1.65], 10.2,
              align=[PP_ALIGN.CENTER, PP_ALIGN.CENTER, PP_ALIGN.CENTER], highlight_rows={2},
              highlight_fill=TEAL_LIGHT)
    text(s, 7.3, 6.35, 5.35, 0.55,
         "31→127 波动 < 0.0004；127 的 +0.000063 与 3 seeds 极差 0.000045 同量级，\n属噪声，不作为“当前最好”。",
         9.2, False, MUTED, line=1.25)
    notes(s, "训练协议是固定 5 折、目标编码只在训练折 fit。学习曲线用 1000 轮 + early stopping 50，"
             "best_iter 489，valid AUC@best 0.96032，末轮 train-valid 差距约 0.0095，全程没有发散。"
             "超参只做了轻量敏感性：num_leaves 31/63/127，波动 <0.0004，默认 63 已近平台；"
             "127 的 +0.000063 在 seed 噪声内，不称为当前最好。")
    return s


def add_slide7(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "06 / 主结果", "主结果：OOF AUC 0.9602，概率校准优秀", 7)

    kpis = [
        ("0.960191", "OOF AUC（5 折均值）", TEAL),
        ("0.960186", "pooled AUC（OOF 拼接）", NAVY),
        ("0.002479", "ECE（10 箱）", GREEN),
        ("0.225458", "LogLoss", AMBER),
    ]
    x = M
    for num, label, color in kpis:
        card(s, x, 1.52, 2.92, 1.05, WHITE, BORDER)
        text(s, x + 0.15, 1.6, 2.6, 0.42, num, 21, True, color)
        text(s, x + 0.15, 2.04, 2.6, 0.25, label, 9.0, False, MUTED)
        x += 3.05

    pic(s, FIG / "p3_calibration.png", M, 2.82, 3.85, 3.05, border=True)
    text(s, M, 5.92, 3.85, 0.27, "校准：预测概率 vs 实际频率", 9.0, False, MUTED, PP_ALIGN.CENTER)
    pic(s, FIG / "p3_confusion.png", 4.65, 2.82, 2.95, 3.05, border=True)
    text(s, 4.65, 5.92, 2.95, 0.27, "混淆矩阵 @ 0.5（仅描述）", 9.0, False, MUTED, PP_ALIGN.CENTER)

    # right bullets
    card(s, 7.95, 2.8, 4.75, 3.35, LIGHT, BORDER, TEAL)
    text(s, 8.15, 2.95, 4.3, 0.28, "怎么读这组数", 10.5, True, TEAL)
    rows = [
        ("95% CI", "[0.959715, 0.960605]"),
        ("阈值 0.5", "Prec 0.9272 · Rec 0.9051"),
        ("F1", "0.9160"),
        ("提交", "299,844 行概率，无 NaN"),
    ]
    yy = 3.32
    for k, v in rows:
        text(s, 8.15, yy, 1.0, 0.26, k, 9.5, True, MUTED)
        text(s, 9.15, yy, 3.4, 0.26, v, 10.2, False, NAVY)
        yy += 0.44
    text(s, 8.15, 5.2, 4.35, 0.75,
         "AUC 是排序指标；阈值指标只说明 0.5 下的行为。\nECE 0.25% → 概率可直接提交。",
         9.2, False, SLATE, line=1.3)
    text(s, M, 6.45, CW, 0.24,
         "最终交付 run_id=P2_final_lgbm_CT_1031_s42_final ｜ 5 折测试集平均 ｜ 299,844 行概率",
         9.0, False, MUTED)
    notes(s, "主结果：5 折均值 0.960191，pooled OOF 0.960186，两者只是算法口径不同；"
             "bootstrap 95% CI 0.959715–0.960605。概率质量：LogLoss 0.2255、Brier 0.0604、ECE 0.0025，"
             "校准很好，不需要 Platt/isotonic。阈值 0.5 下的 Acc/Precision/Recall 只作描述，不用于选模。"
             "submission 299,844 行概率、无 NaN。")
    return s


def add_slide8(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "07 / 稳定性", "3 seeds × 5 folds：换 seed 几乎不动，换折才是主因", 8)

    kpis = [
        ("0.960184", "seed 均值 AUC", TEAL),
        ("0.000023", "seed 间 std", AMBER),
        ("0.000631", "平均折内 std（约 27×）", RED),
    ]
    x = M
    for num, label, color in kpis:
        card(s, x, 1.5, 3.92, 0.9, WHITE, BORDER, color)
        text(s, x + 0.16, 1.58, 3.55, 0.4, num, 20, True, color)
        text(s, x + 0.16, 2.0, 3.55, 0.24, label, 9, False, MUTED)
        x += 4.08

    pic(s, FIG / "p3_seed_variance.png", M, 2.72, 7.35, 2.95, border=True)
    text(s, M, 5.72, 7.35, 0.28, "每 seed 的 OOF AUC；虚线为 3-seed 均值", 9.0, False, MUTED, PP_ALIGN.CENTER)

    data = [
        ["seed", "OOF AUC", "fold std", "Δ"],
        ["42", "0.960191", "0.000583", "≈0"],
        ["2026", "0.960159", "0.000650", "−0.000032"],
        ["7", "0.960203", "0.000660", "+0.000012"],
        ["均值", "0.960184", "0.000631", "−0.000007"],
    ]
    add_table(s, 8.15, 2.72, 4.55, 2.35, data, [0.8, 1.2, 1.15, 1.15], 9.8,
              align=[PP_ALIGN.CENTER] * 4, highlight_rows={4}, highlight_fill=TEAL_LIGHT)
    text(s, 8.15, 5.25, 4.55, 1.0,
         "seed 极差 0.000045，远小于折间 std。\n"
         "C+T 的 +0.001362 仍约为 seed 极差的 30 倍，\n但 3 seeds 只是方差估计，不是正式显著性检验。",
         9.2, False, SLATE, line=1.3)
    notes(s, "稳定性是 3 seeds × 固定 5 折，只换 random_state：0.960191 / 0.960159 / 0.960203，"
             "均值 0.960184，seed 间 std 只有 0.000023，极差 0.000045；而平均折内 std 是 0.000631，约 27 倍。"
             "所以换 seed 几乎不动，换折才是主要不确定性；C+T 的 +0.00136 约为 seed 极差 30 倍，"
             "但在该划分下仍不能称统计显著。")
    return s


def add_slide9(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "08 / 误差分析", "错得最多的不是差评，而是决策边界样本", 9)

    pic(s, FIG / "p3_badcase.png", M, 1.55, 6.55, 4.2, border=True)
    text(s, M, 5.82, 6.55, 0.28, "分群错误率（0.5 阈值，仅用于定位）", 9.0, False, MUTED, PP_ALIGN.CENTER)

    # right stats
    stats = [("7.36%", "0.5 阈值错误率（51,490 条）", RED),
             ("14,497", "高置信 FP：y=0 且 pred≥0.8", AMBER),
             ("16,146", "高置信 FN：y=1 且 pred≤0.2", TEAL)]
    yy = 1.55
    for num, label, color in stats:
        text(s, 7.55, yy, 1.85, 0.42, num, 22, True, color)
        text(s, 9.35, yy + 0.04, 3.35, 0.3, label, 9.5, False, SLATE)
        yy += 0.62

    card(s, 7.55, 3.55, 5.15, 1.12, AMBER_LIGHT, RGBColor(0xF5, 0xC9, 0x8A), AMBER)
    text(s, 7.72, 3.66, 4.8, 0.26, "错误率最高的群：Online boarding = 4", 10.5, True, AMBER)
    text(s, 7.72, 3.96, 4.8, 0.55, "12.0%  （错误 25,310 / 该群 210,950）", 12.5, True, NAVY)

    text(s, 7.55, 4.85, 5.15, 1.4,
         "• Business 以 FP 为主（16,485）：模型偏乐观\n"
         "• Eco 以 FN 为主（20,923）：模型偏悲观\n"
         "• Online boarding=4 处于满意/不满意边界，\n  标签本身可能带主观随机性",
         9.6, False, SLATE, line=1.35)

    rect(s, M, 6.35, CW, 0.5, NAVY)
    text(s, M + 0.2, 6.43, CW - 0.4, 0.3,
         "这些误判样本不是“差评”，而是评分组合卡在边界；按舱位做阈值/校准比继续调参更实际。",
         11.2, True, WHITE)
    notes(s, "误差分析先看总体：0.5 阈值错误 51,490 条，7.36%；高置信 FP 14,497、FN 16,146。"
             "最有信息量的发现是错误率最高的不是差评，而是 Online boarding=4 的边界样本，12.0% 错误率。"
             "Business 以 FP 为主、Eco 以 FN 为主，说明分舱位做阈值/校准会有收益。")
    return s


def add_slide10(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "09 / 可解释性", "SHAP：模型真正在用哪些信号", 10)

    pic(s, FIG / "p3_shap_importance.png", M, 1.55, 4.35, 4.75, border=True)
    text(s, M, 6.35, 4.35, 0.28, "mean |SHAP|，20,000 行采样，OOF 目标编码", 9.0, False, MUTED, PP_ALIGN.CENTER)

    data = [
        ["排名", "特征", "SHAP 占比"],
        ["1", "Online boarding", "20.24%"],
        ["2", "Type of Travel", "14.69%"],
        ["3", "Inflight wifi service", "12.39%"],
        ["4", "Customer Type", "8.44%"],
        ["5", "Class", "8.30%"],
    ]
    add_table(s, 5.35, 1.55, 4.25, 2.55, data, [0.65, 2.45, 1.15], 10.2,
              align=[PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.CENTER], highlight_rows={1},
              highlight_fill=TEAL_LIGHT)

    card(s, 9.85, 1.55, 2.85, 2.55, LIGHT, BORDER, TEAL)
    text(s, 10.02, 1.7, 2.5, 0.28, "gain vs SHAP", 10, True, TEAL)
    text(s, 10.02, 2.02, 2.55, 1.95,
         "Online boarding\ngain 42.1% → SHAP 20.2%\n\n"
         "Flight Distance__te\ngain 1.06% → SHAP 5.82%",
         9.2, False, SLATE, line=1.3)
    text(s, 5.35, 4.35, 7.35, 1.55,
         "• gain 回答“模型在哪里分裂”，SHAP 回答“特征把预测推了多少”，两者不能互相替代。\n"
         "• TE 的 SHAP 贡献 5.82% 高于 gain 1.06%：它改变很多行的分数，但分裂次数少。\n"
         "• T 组用逐折 OOF 编码，避免自泄漏；SHAP 不是因果，不能回答“提高 WiFi 会不会更满意”。",
         9.8, False, SLATE, line=1.4)
    notes(s, "可解释性用 LightGBM 的 pred_contrib 算 TreeSHAP，采样 2 万行，T 组用逐折 OOF 编码避免自泄漏。"
             "前五：Online boarding 20.24%、Type of Travel 14.69%、WiFi 12.39%、Customer Type 8.44%、Class 8.30%。"
             "重点对比 gain 和 SHAP：Online boarding gain 42% 但 SHAP 20%，TE gain 只有 1% 但 SHAP 5.8%，"
             "说明两者回答不同问题。SHAP 不是因果。")
    return s


def add_slide11(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "10 / 结论", "结论、局限与下一步", 11)

    cards = [
        ("主结果", "OOF 0.960191（5 折均值）/ 0.960186（pooled），95% CI 0.9597–0.9606", TEAL),
        ("稳定性", "3 seeds 均值 0.960184，seed std 0.000023；折间 std 0.000631 是主因", AMBER),
        ("概率质量", "ECE 0.002479，LogLoss 0.2255；概率可直接提交", GREEN),
        ("误差", "错误集中在 Online boarding=4 边界样本；Business FP / Eco FN", RED),
    ]
    x, y = M, 1.6
    for i, (h, b, color) in enumerate(cards):
        card(s, x, y, 3.05, 1.55, WHITE, BORDER, color)
        text(s, x + 0.18, y + 0.15, 2.7, 0.28, h, 11, True, color)
        text(s, x + 0.18, y + 0.5, 2.68, 0.9, b, 9.4, False, SLATE, line=1.3)
        x += 3.15
        if i == 1:
            x, y = M, 3.3

    card(s, 9.65, 1.6, 3.05, 3.25, LIGHT, BORDER, NAVY)
    text(s, 9.82, 1.75, 2.7, 0.28, "局限", 10.5, True, NAVY)
    text(s, 9.82, 2.08, 2.72, 2.6,
         "• 公开镜像数据，不是官方原文件\n"
         "• 没有官方 Public LB 提交\n"
         "• 3 seeds 不是正式显著性检验\n"
         "• 目标编码对分布漂移敏感\n"
         "• SHAP 为 2 万行采样，非因果",
         9.0, False, SLATE, line=1.35)

    # bottom: outlook + innovation
    card(s, M, 5.15, 7.9, 1.15, TEAL_LIGHT, RGBColor(0xB8, 0xDD, 0xE0), TEAL)
    text(s, M + 0.2, 5.27, 7.5, 0.26, "下一步", 10, True, TEAL)
    text(s, M + 0.2, 5.56, 7.5, 0.62,
         "官方数据补 LB · CatBoost/XGBoost 横向对比 · 分舱位校准与业务阈值 · 部署后概率校准监控",
         9.8, False, SLATE, line=1.3)
    rect(s, 8.7, 5.15, 4.0, 1.15, NAVY)
    text(s, 8.9, 5.28, 3.6, 0.26, "统一口径 · 创新点", 9.5, True, AMBER)
    text(s, 8.9, 5.58, 3.65, 0.55, "严格防泄漏协议\n+ 精简化特征 + 分析型交付", 10.5, True, WHITE, line=1.25)

    text(s, M, 6.5, CW, 0.3,
         "结论：0.96 的天花板已被卷平；本项目的价值在于把“为什么可信”讲清楚，而不是再堆一个模型。",
         10.5, False, MUTED)
    notes(s, "结论四条：主结果、稳定性、概率质量、误差定位。局限要主动讲：镜像数据、没有官方 LB、"
             "3 seeds 只是方差估计、TE 对漂移敏感、SHAP 采样且非因果。下一步是官方数据补 LB、"
             "CatBoost/XGBoost 对比、分舱位校准。最后统一口径：创新点是严格防泄漏 + 精简化特征 + 分析型交付。")
    return s


def add_slide12(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "11 / 团队", "分工、交接与致谢", 12)

    roles = [
        ("P1", "数据与特征", ["数据字典 / EDA / 划分", "特征 A/B/C/T 与消融", "特征字典 / 泄漏排查"], TEAL),
        ("P2", "模型与训练", ["LightGBM / LR 对照", "调参敏感性与学习曲线", "最终模型 / submission"], AMBER),
        ("P3", "评估与交付", ["评估协议 / 校准", "3 seeds / SHAP / badcase", "PPT 统稿与验收"], NAVY),
    ]
    x = M
    for code, name, items, color in roles:
        card(s, x, 1.55, 3.95, 2.75, WHITE, BORDER, color)
        text(s, x + 0.2, 1.7, 1.2, 0.3, code, 11, True, color)
        text(s, x + 0.2, 2.0, 3.55, 0.32, name, 14, True, NAVY)
        yy = 2.45
        for it in items:
            text(s, x + 0.2, yy, 3.55, 0.28, f"·  {it}", 9.4, False, SLATE)
            yy += 0.34
        text(s, x + 0.2, 3.85, 3.55, 0.26, "✓ 已完成", 9.5, True, GREEN)
        x += 4.08

    # handoff timeline
    text(s, M, 4.6, 3.0, 0.28, "关键交接 H1–H4", 10.5, True, TEAL)
    chips = [("H1", "冻结划分"), ("H2", "模型接口"), ("H3", "最终 submission"), ("H4", "三份报告")]
    x = M
    for code, label in chips:
        rect(s, x, 4.95, 2.75, 0.62, LIGHT, BORDER, rounded=True, radius=0.15)
        text(s, x + 0.15, 5.03, 0.5, 0.26, code, 10, True, TEAL)
        text(s, x + 0.62, 5.04, 1.95, 0.26, label, 9.5, False, NAVY)
        text(s, x + 2.45, 5.03, 0.25, 0.26, "✓", 10, True, GREEN)
        x += 2.9

    rect(s, M, 5.95, CW, 0.72, NAVY)
    text(s, M + 0.25, 6.03, 6.2, 0.3, "感谢课程方、Kaggle 与公开镜像作者", 12, True, WHITE)
    text(s, SW - M - 3.0, 6.03, 2.75, 0.35, "Q & A", 15.5, True, AMBER, PP_ALIGN.RIGHT)
    text(s, M, 6.78, CW, 0.24,
         "工程验收：119 pytest · verify 34/34 · 全链路可重建 ｜ 剩余：PPT 导出 · 双面打印 · 预答辩彩排",
         9.0, False, MUTED)
    notes(s, "最后一页：P1 数据与特征、P2 模型与训练、P3 评估与交付，三条线都已完成；"
             "H1 冻结划分、H2 模型接口、H3 最终 submission、H4 三份报告全部闭环。"
             "感谢课程方、Kaggle 和公开镜像作者；剩余只有 PPT 导出、打印和彩排。进入 Q&A。")
    return s


def add_table(slide, x, y, w, h, data, col_widths, font_size=10.5, align=None,
              highlight_rows=(), highlight_fill=TEAL_LIGHT):
    rows, cols = len(data), len(data[0])
    shape = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    table = shape.table
    table.first_row = True
    table.horz_banding = False
    total_w = sum(col_widths)
    col_widths = [cw * w / total_w for cw in col_widths]
    for i, cw in enumerate(col_widths):
        table.columns[i].width = Inches(cw)
    row_h = h / rows
    for r in range(rows):
        table.rows[r].height = Inches(row_h)
        for c in range(cols):
            cell = table.cell(r, c)
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = NAVY
            elif r in highlight_rows:
                cell.fill.fore_color.rgb = highlight_fill
            else:
                cell.fill.fore_color.rgb = WHITE if r % 2 else LIGHT
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = (align[c] if align else PP_ALIGN.CENTER)
            run = p.add_run()
            run.text = str(data[r][c])
            _font(run, font_size, bold=(r == 0 or r in highlight_rows),
                  color=WHITE if r == 0 else INK)
    return table


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    prs.core_properties.title = "航空乘客满意度预测 · Kaggle S6E10"
    prs.core_properties.author = "P1 / P2 / P3 课程项目组"
    prs.core_properties.subject = "数据、特征、模型与评估全流程"
    prs.core_properties.comments = "由 scripts/make_ppt.py 生成；数字来自 reports/exp_log.csv 与 reports/eval/。"
    add_cover(prs)
    add_slide2(prs)
    add_slide3(prs)
    add_slide4(prs)
    add_slide5(prs)
    add_slide6(prs)
    add_slide7(prs)
    add_slide8(prs)
    add_slide9(prs)
    add_slide10(prs)
    add_slide11(prs)
    add_slide12(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"[ppt] {len(prs.slides)} 页 -> {OUT}")
    print(f"[ppt] 大小 {OUT.stat().st_size / 1024:.0f} KB")


if __name__ == "__main__":
    main()
